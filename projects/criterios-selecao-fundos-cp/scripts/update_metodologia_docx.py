#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pathlib import Path
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

ROOT = Path("projects/criterios-selecao-fundos-cp")
DOCX = ROOT / "Critérios de Seleção de Fundos CP_v4.docx"

def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)

def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        yield from iter_table_paragraphs(table)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                yield p
            for table in container.tables:
                yield from iter_table_paragraphs(table)

def replace_paragraph_text(paragraph, new_text):
    if paragraph.text == new_text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)

def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    if paragraph.style is not None:
        created.style = paragraph.style
    created.add_run(text)
    return created

def apply_rules(paragraph):
    original = paragraph.text
    text = original
    low = text.lower()

    # Pesos explícitos por métrica.
    if "hit rate mensal" in low or "hit mensal" in low:
        text = re.sub(r"(?<!\d)40\s*%", "20%", text)
    if "hit rate 12" in low or "hit 12" in low:
        text = re.sub(r"(?<!\d)40\s*%", "60%", text)

    # Frases que trazem os três pesos juntos.
    patterns = [
        (
            r"(?i)(hit rates?\s*)?mensal\s*40\s*%\s*[,;/|-]*\s*6\s*meses\s*20\s*%\s*(?:e|[,;/|-])\s*12\s*meses\s*40\s*%",
            "Hit rates mensal 20%, 6 meses 20% e 12 meses 60%"
        ),
        (
            r"(?i)mensal\s*\(\s*40\s*%\s*\)\s*[,;/|-]*\s*6\s*meses\s*\(\s*20\s*%\s*\)\s*(?:e|[,;/|-])\s*12\s*meses\s*\(\s*40\s*%\s*\)",
            "mensal (20%), 6 meses (20%) e 12 meses (60%)"
        ),
    ]
    for pattern, repl in patterns:
        text = re.sub(pattern, repl, text)

    # Se o parágrafo descreve consistência/hit rate, a janela não é mais 36m.
    low_after = text.lower()
    if ("consist" in low_after or "hit rate" in low_after or "hit mensal" in low_after) and "36 meses" in low_after:
        text = re.sub(r"(?i)36\s+meses", "72 meses", text)

    # Fórmulas descritivas antigas.
    if "excesso positivo" in low_after and ("dividid" in low_after or "meses" in low_after):
        text = re.sub(r"(?i)dividid[oa]s?\s+por\s+36", "divididos por 72", text)

    if text != original:
        replace_paragraph_text(paragraph, text)
        return True, original, text
    return False, original, text

def main():
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)

    doc = Document(DOCX)
    changes = []
    paragraphs = list(iter_paragraphs(doc))

    for p in paragraphs:
        changed, before, after = apply_rules(p)
        if changed:
            changes.append((before, after))

    # Releitura após alterações para assegurar uma descrição inequívoca.
    paragraphs = list(iter_paragraphs(doc))
    full_text = "\n".join(p.text for p in paragraphs)

    explanatory = (
        "Os três hit rates são calculados sobre a mesma janela comum de 72 meses: "
        "mensal em 72 observações, 6 meses em 67 janelas móveis e 12 meses em "
        "61 janelas móveis. Dentro do pilar de consistência, os pesos são 20%, "
        "20% e 60%, respectivamente."
    )

    if "61 janelas móveis" not in full_text:
        anchor = None
        for p in paragraphs:
            low = p.text.lower()
            if "consist" in low and ("hit" in low or "25%" in low):
                anchor = p
                break
        if anchor is None:
            for p in paragraphs:
                if "hit rate" in p.text.lower():
                    anchor = p
                    break
        if anchor is None:
            anchor = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
        insert_after(anchor, explanatory)
        changes.append(("<nota explicativa ausente>", explanatory))

    # Pós-condições metodológicas.
    final_text = "\n".join(p.text for p in iter_paragraphs(doc))
    required = [
        "72 meses",
        "67 janelas móveis",
        "61 janelas móveis",
        "20%",
        "60%",
    ]
    missing = [x for x in required if x not in final_text]
    if missing:
        raise RuntimeError("DOCX sem elementos metodológicos obrigatórios: " + ", ".join(missing))

    doc.save(DOCX)

    print(f"[DOCX] {len(changes)} alteração(ões) aplicada(s).")
    for before, after in changes:
        print("[DOCX] ANTES:", before[:500])
        print("[DOCX] DEPOIS:", after[:500])

    # Trechos finais relevantes para auditoria no log.
    print("[DOCX] TRECHOS FINAIS RELEVANTES:")
    for p in iter_paragraphs(Document(DOCX)):
        low = p.text.lower()
        if "consist" in low or "hit rate" in low or "61 janelas" in low:
            print(" -", p.text[:1000])

if __name__ == "__main__":
    main()
